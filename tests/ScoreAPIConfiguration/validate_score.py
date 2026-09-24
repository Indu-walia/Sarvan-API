import sys
import time
import warnings
import requests
import openpyxl
from pathlib import Path
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from openpyxl.styles import PatternFill, Font

warnings.filterwarnings('ignore')

MAX_WORKERS = 3
SAVE_EVERY = 1000
REQUEST_DELAY = 0.15  # seconds, paced per request to avoid overloading the server under sustained load

SCRIPT_DIR = Path(__file__).resolve().parent
DATA_FILE  = SCRIPT_DIR / 'ScoreAPIConfigurationData.xlsx'

API_URL = 'https://p9authwave.mox2.net.in/p9/reprocesslog.ashx'

# Sheet name → API key mapping
SHEET_KEYS = {
    'Client_Key1_85BE':  '85BE-52EC-6E0A-52F2-8F17-E377-3FFC-722F',
    'BFL_Key2_2DAC':     '2DAC-DC08-8B5F-50D3-341E-A4FE-2728-7CDB',
    'Default_Key3_F08C': 'F08C-22B9-5A7E-974F-74E3-3A6F-8316-F33E',
}

# Language name → API lang param
LANG_MAP = {
    'assamese': 'assamese', 'bengali': 'bengali',  'gujarati': 'gujarati',
    'hindi':    'hindi',    'kannada': 'kannada',   'malayalam': 'malayalam',
    'marathi':  'marathi',  'oriya':   'oriya',     'punjabi':  'punjabi',
    'tamil':    'tamil',    'telugu':  'telugu',     'urdu':     'urdu',
}

GREEN = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
RED   = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
BOLD  = Font(bold=True)

# Column positions (1-based)
COL_LANG         = 2
COL_QUAL         = 3
COL_SRC          = 4
COL_EXP_SCORE    = 5
COL_ACT_SCORE    = 6
COL_EXP_OUTQUAL  = 7
COL_OUTQUAL      = 8
COL_RESULT       = 9
COL_OUTQUAL_RES  = 10
COL_OVERALL      = 11

_session = requests.Session()


def call_api(text, quality, language, api_key):
    lang_param = LANG_MAP.get(language.lower(), language.lower())
    payload = {
        'key':           api_key,
        'data':          [{'text': text, 'qual': str(quality), 'op': '0'}],
        'InputLanguage': 'english',
        'lang':          [lang_param],
    }
    time.sleep(REQUEST_DELAY)
    resp = _session.post(API_URL, json=payload, timeout=30, verify=False)
    resp.raise_for_status()
    return resp.json()[0]


def score_result(actual, expected):
    """Literal comparison: Pass if actual score >= the Expected Score given, for every quality."""
    if expected is None:
        return 'N/A'
    return 'Pass' if actual >= expected else 'Fail'


def outqual_result(actual_outqual, requested_quality):
    """Expected OutQual is simply the Requested Quality itself. Pass if actual OutQual matches it."""
    try:
        return 'Pass' if int(actual_outqual) == int(requested_quality) else 'Fail'
    except (TypeError, ValueError):
        return 'Fail'


def overall_result(score_res, outqual_res):
    if outqual_res == 'N/A':
        return score_res
    if score_res == 'N/A':
        return outqual_res
    return 'Pass' if score_res == 'Pass' and outqual_res == 'Pass' else 'Fail'


def _fetch(row, lang, qual, src_text, exp_score, api_key):
    try:
        data = call_api(src_text, qual, lang, api_key)
        return row, exp_score, data, None
    except Exception as e:
        return row, exp_score, None, e


def process_sheet(ws, api_key, save_cb):
    jobs = []
    for row in range(2, ws.max_row + 1):
        lang      = ws.cell(row, COL_LANG).value
        qual      = ws.cell(row, COL_QUAL).value
        src_text  = ws.cell(row, COL_SRC).value
        exp_score = ws.cell(row, COL_EXP_SCORE).value

        if not lang or not qual or not src_text:
            continue
        jobs.append((row, lang, qual, src_text, exp_score))

    total = len(jobs)
    done = 0

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        futures = [pool.submit(_fetch, row, lang, qual, src_text, exp_score, api_key)
                   for row, lang, qual, src_text, exp_score in jobs]

        for future in as_completed(futures):
            row, exp_score, data, err = future.result()
            done += 1

            if err is not None:
                for col in (COL_ACT_SCORE, COL_EXP_OUTQUAL, COL_OUTQUAL, COL_RESULT, COL_OUTQUAL_RES, COL_OVERALL):
                    ws.cell(row, col, f'ERROR: {err}')
            else:
                qual         = ws.cell(row, COL_QUAL).value
                actual_score = data.get('score') or 0.0
                out_qual     = data.get('outQual', '')
                result       = score_result(round(actual_score, 6), exp_score)
                outqual_res  = outqual_result(out_qual, qual)
                overall      = overall_result(result, outqual_res)

                ws.cell(row, COL_ACT_SCORE,   round(actual_score, 6))
                ws.cell(row, COL_EXP_OUTQUAL, qual)
                ws.cell(row, COL_OUTQUAL,     out_qual)
                ws.cell(row, COL_RESULT,      result)

                for col, val in ((COL_OUTQUAL_RES, outqual_res), (COL_OVERALL, overall)):
                    cell = ws.cell(row, col, val)
                    if val == 'Pass':
                        cell.fill = GREEN
                    elif val == 'Fail':
                        cell.fill = RED

            if done % 100 == 0 or done == total:
                print(f'  {done}/{total} rows done', flush=True)
            if done % SAVE_EVERY == 0:
                save_cb()


def main():
    wb = openpyxl.load_workbook(DATA_FILE)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    out_file  = str(SCRIPT_DIR / f'ScoreAPIConfigurationData_Results_{timestamp}.xlsx')

    def checkpoint():
        wb.save(out_file)
        print(f'  [checkpoint saved: {out_file}]', flush=True)

    for sheet_name, api_key in SHEET_KEYS.items():
        if sheet_name not in wb.sheetnames:
            print(f'Sheet not found: {sheet_name} — skipping')
            continue
        print(f'\n=== {sheet_name} ===')
        process_sheet(wb[sheet_name], api_key, checkpoint)
        checkpoint()

    wb.save(out_file)
    print(f'\nSaved: {out_file}')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
