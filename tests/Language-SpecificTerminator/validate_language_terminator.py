import sys
import warnings
import requests
import openpyxl
from datetime import datetime
from openpyxl.styles import PatternFill, Font

warnings.filterwarnings('ignore')

API_URL = 'https://p9authwave.mox2.net.in/p9/reprocesslog.ashx'
API_KEY = '352B-6597-8C4E-53BD-AEE8-5BCA-7E83-3457'

LANGUAGES = [
    ('hindi',               '।', 'Hindi - Devanagari Danda'),
    ('sanskrit',            '।', 'Sanskrit - Devanagari Danda'),
    ('nepali',              '।', 'Nepali - Devanagari Danda'),
    ('assamese',            '।', 'Assamese - Devanagari Danda'),
    ('punjabi',             '।', 'Punjabi - Devanagari Danda'),
    ('oriya',               '।', 'Oriya - Devanagari Danda'),
    ('bengali',             '।', 'Bengali - Devanagari Danda'),
    ('chinese_traditional', '。', 'Chinese Traditional - Ideographic Full Stop'),
    ('chinese_simplified',  '。', 'Chinese Simplified - Ideographic Full Stop'),
    ('japanese',            '。', 'Japanese - Ideographic Full Stop'),
    ('urdu',                '۔', 'Urdu - Arabic Full Stop'),
    ('arabic',              '۔', 'Arabic - Arabic Full Stop'),
    ('korean',              '.',  'Korean - Standard Full Stop'),
    ('tamil',               '.',  'Tamil - Standard Full Stop'),
]

GREEN = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
RED   = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
BOLD  = Font(bold=True)

INPUT_FILE = 'LanguageTerminatorData.xlsx'


def call_api(text, language):
    payload = {
        'key': API_KEY,
        'data': [{'text': text, 'qual': '4', 'op': '0'}],
        'InputLanguage': 'english',
        'lang': [language]
    }
    resp = requests.post(API_URL, json=payload, timeout=30, verify=False)
    resp.raise_for_status()
    return resp.json()


def check_terminator(output_text, exp_terminator):
    if not exp_terminator:
        return 'Pass', ''
    if exp_terminator not in output_text:
        return 'Fail', f'Expected "{exp_terminator}" not found in output'
    if exp_terminator + exp_terminator in output_text:
        return 'Fail', f'Double terminator "{exp_terminator}{exp_terminator}" found in output'
    return 'Pass', ''


def main():
    in_wb = openpyxl.load_workbook(INPUT_FILE, read_only=True)
    in_ws = in_wb.active

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active

    headers = [
        'Sno', 'Description', 'SourceText',
        'Language', 'ExpectedTerminator',
        'ActualOutputText',
        'Terminator_Result', 'Overall_Result', 'FailReason'
    ]
    for col, h in enumerate(headers, 1):
        cell = out_ws.cell(1, col, h)
        cell.font = BOLD

    out_row = 2
    sno     = 1

    for row_data in in_ws.iter_rows(min_row=2, values_only=True):
        source_text = (row_data[2] or '').strip()
        if not source_text:
            continue

        print(f'\nSource: {source_text[:70]}')

        for lang, exp_terminator, lang_desc in LANGUAGES:
            print(f'  [{lang}]...', end=' ', flush=True)

            try:
                all_data    = call_api(source_text, lang)
                output_text = ' '.join(
                    d.get('outputText', '') for d in all_data if isinstance(d, dict)
                )
                result, reason = check_terminator(output_text, exp_terminator)
            except requests.HTTPError as e:
                output_text = ''
                result  = 'Fail'
                reason  = f'HTTP {e.response.status_code}: {e.response.text[:80]}'
            except Exception as e:
                output_text = ''
                result  = 'Fail'
                reason  = f'ERROR: {e}'

            print(result)

            row_vals = [
                sno, lang_desc, source_text,
                lang, exp_terminator,
                output_text,
                result, result, reason
            ]
            for col, val in enumerate(row_vals, 1):
                cell = out_ws.cell(out_row, col, val)
                if val in ('Pass', 'Fail'):
                    cell.fill = GREEN if val == 'Pass' else RED

            out_row += 1
            sno     += 1

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    out_file  = f'LanguageTerminatorData_Results_{timestamp}.xlsx'
    out_wb.save(out_file)
    print(f'\nSaved: {out_file}')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
