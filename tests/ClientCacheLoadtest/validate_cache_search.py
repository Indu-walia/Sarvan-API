import sys
import time
import warnings
import requests
import openpyxl
from datetime import datetime
from openpyxl.styles import PatternFill, Font

warnings.filterwarnings('ignore')

API_URL  = 'https://clienttm.moxwave.com/moxlocalization/moxtm/GetGridRecords'
AUTH_KEY = '8E18-265B-E3FC-9A1A-8EF9-282A-F395-3FDA'
LANGUAGE = 'Hindi'
QUALITY  = 6

GREEN = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
RED   = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
BOLD  = Font(bold=True)


def search(text):
    payload = {
        'AuthKey': AUTH_KEY,
        'Language': LANGUAGE,
        'Quality': QUALITY,
        'SearchText': text,
        'FieldName': 'source_text',
        'PageIndex': 0,
        'RowCount': 10
    }
    start = time.perf_counter()
    resp = requests.post(API_URL, json=payload, timeout=30, verify=False)
    elapsed = time.perf_counter() - start
    resp.raise_for_status()
    return resp.json(), elapsed


def main():
    wb = openpyxl.load_workbook('CCSearchloadtest.xlsx')
    ws = wb.active

    headers = ['Searched', 'TimeTaken_ms', 'MatchCount', 'Target_text']
    start_col = ws.max_column + 1
    for idx, h in enumerate(headers):
        ws.cell(1, start_col + idx, h).font = BOLD

    for row in range(2, ws.max_row + 1):
        src_text = ws.cell(row, 1).value
        if not src_text or not str(src_text).strip():
            continue

        print(f'Row {row}...', end=' ', flush=True)

        try:
            data, elapsed = search(str(src_text))
            total = data.get('total', 0)
            records = data.get('records', [])
            searched = 'Yes' if total > 0 and records else 'No'
            target = records[0].get('Target_text', '') if records else ''
        except Exception as e:
            searched, elapsed, total, target = f'ERROR: {e}', '', 0, ''

        elapsed_ms = round(elapsed * 1000, 2) if isinstance(elapsed, float) else elapsed

        ws.cell(row, start_col, searched)
        ws.cell(row, start_col).fill = GREEN if searched == 'Yes' else RED
        ws.cell(row, start_col + 1, elapsed_ms)
        ws.cell(row, start_col + 2, total)
        ws.cell(row, start_col + 3, target)

        print(f'Searched={searched}, Time={elapsed_ms}ms, Matches={total}')

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    out_file = f'CCSearchloadtest_Results_{timestamp}.xlsx'
    wb.save(out_file)
    print(f'\nSaved: {out_file}')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
